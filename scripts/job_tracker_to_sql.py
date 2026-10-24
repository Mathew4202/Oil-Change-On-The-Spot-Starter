import openpyxl, re, json, sys
src, out = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(src, data_only=True)

MAKES = {'nissan':'Nissan','nissans':'Nissan','hyundai':'Hyundai','hyuandai':'Hyundai','honda':'Honda','toyota':'Toyota',
 'chevy':'Chevrolet','chevrolet':'Chevrolet','gmc':'GMC','jeep':'Jeep','ford':'Ford','mazda':'Mazda','subaru':'Subaru',
 'kia':'Kia','vw':'Volkswagen','volkswagen':'Volkswagen','bmw':'BMW','mercedes':'Mercedes-Benz','mitsubishi':'Mitsubishi',
 'mistubishi':'Mitsubishi','mitsubish':'Mitsubishi','chrysler':'Chrysler','dodge':'Dodge','scion':'Scion','acura':'Acura','ram':'Ram'}
MODEL_FIX = {'foresta':'Forester','forester':'Forester','crv':'CR-V','cruse':'Cruze','qashqui':'Qashqai','altas':'Atlas','orlando':'Orlando',
 'cx-30':'CX-30','ilx':'ILX','rvr':'RVR','im':'iM','f150':'F-150','x3':'X3','c300':'C300'}
MODEL_MAKE = {'elantra':'Hyundai','cruze':'Chevrolet','civic':'Honda','corolla':'Toyota'}

def q(v):
    if v is None or v == '': return 'null'
    if isinstance(v,(int,float)): return repr(round(float(v),2))
    return "'" + str(v).replace("'","''").strip() + "'"

def num(v):
    return float(v) if isinstance(v,(int,float)) else None

def parse_vehicle(t):
    d = {'year':None,'make':None,'model':None,'engine':None,'grade':None,'cap':None,'filter':None}
    if not t: return d
    s = ' '.join(str(t).split())
    low = s.lower()
    m = re.search(r'\b(19[89]\d|20[0-3]\d)\b', s); d['year'] = int(m.group(1)) if m else None
    words = re.findall(r"[A-Za-z0-9\-]+", s.split('(')[0])
    if not any(w.lower() in MAKES for w in words):
        for w in words:
            if w.lower() in MODEL_MAKE: words.insert(words.index(w), MODEL_MAKE[w.lower()]); break
    for i,w in enumerate(words):
        if w.lower() in MAKES:
            d['make'] = MAKES[w.lower()]
            rest = [x for x in words[i+1:i+4] if not re.match(r'^\d', x) or re.match(r'^\d+$', x) and len(x)<=2 or x in ('4Runner','3','6','200','1500','300','500')]
            mw = []
            for x in words[i+1:i+4]:
                xl = x.lower()
                if re.match(r'^\d\.\d', xl) or xl in ('oil','engine','v6','v8','my','hybrid','lt','le','lx','se','sv','ex','gt','xle','essential','platinum','limited','classic','awd','eco','i4'): break
                if re.match(r'^(19|20)\d\d$', xl): continue
                if re.match(r'^\d+$', xl) and mw: break
                if xl in ('lost','cabin','customer','turbo','all','terrain'): break
                mw.append(MODEL_FIX.get(xl, x if not x.islower() else x.capitalize()))
                if len(mw) == 2 or xl in ('cruze','civic','corolla','camry','accord','elantra','tucson','sonata','rogue','altima','maxima','pilot','kona','soul','rio','malibu','f150','tundra','pathfinder','equinox','journey','fusion','escape','outback','impreza','venue','kicks','sentra','terrain','highlander','rav4','trax','ilx','tiguan','atlas','cx-30','lancer','rvr','outlander','4runner','sierra','x3','c300','qashqai','sorento','forester','foresta','crv','orlando','accent','santa','grand','eclipse','genesis','im','200','1500','6','3'):
                    if xl in ('grand','eclipse','santa','genesis') and len(mw)==1: continue
                    break
            d['model'] = ' '.join(mw) or None
            break
    pre = re.split(r'\(|\boil\b|\bengine type\b|\d{1,2}\s*w\s*[-_]?\s*\d', low)[0]
    m = re.search(r'\b(\d(?:\.\d)?)\s*(?:l|litre|liter)\b', pre) or re.search(r'\b(\d\.\d)\b', pre)
    if not m:
        m = re.search(r'engine\s*-\s*(\d\.\d)\s*l', low)
    d['engine'] = (f"{float(m.group(1)):.1f}L") if m and float(m.group(1)) <= 6.8 else None
    g = re.search(r'\b(\d{1,2})\s*w\s*[-_]?\s*(\d{1,2})\b', low)
    if g:
        d['grade'] = f"{g.group(1)}W-{g.group(2)}"
        if 'euro' in low[g.start():g.start()+40] or 'vw50' in low or '229' in low: d['grade'] = 'Euro ' + d['grade']
        c = re.search(r'(\d{1,2}(?:\.\d)?)\s*l\b', low[g.end():g.end()+40])
    else:
        c = None
    if not c:
        c = re.search(r'capacity\s*[-:]?\s*(\d{1,2}(?:\.\d)?)\s*l', low)
    if c:
        v = float(c.group(1))
        if 2.5 <= v <= 9: d['cap'] = v
    f = re.search(r'(?<!cabin )(?<!engine )filter(?:\s*type)?\s*[-–:]+\s*(\d{4,7})', low)
    d['filter'] = f.group(1) if f else None
    return d

def method(n):
    n = (n or '').lower()
    if 'cash' in n: return 'cash'
    if 'transfer' in n: return 'etransfer'
    return None

rows = []
ws = wb['Job Tracker']
for r in ws.iter_rows(min_row=2, values_only=True):
    if not any(c not in (None,'') for c in r[:5]): continue
    date, name, veh, svc, price, oil, filt, gas, other, tcost, profit, notes, split = (list(r)+[None]*13)[:13]
    v = parse_vehicle(veh)
    rows.append(dict(date=date.strftime('%Y-%m-%d'), name=' '.join(str(name or '').split()), raw=' '.join(str(veh or '').split()) or None,
        service=' '.join(str(svc or '').split()), price=num(price), oil=num(oil), filt=num(filt),
        other=(num(gas) or 0)+(num(other) or 0) if (num(gas) is not None or num(other) is not None) else None,
        notes=' '.join(str(notes).split()) if notes and not method(str(notes)) or (notes and len(str(notes))>12) else None,
        method=method(str(notes) if notes else ''), split=' '.join(str(split).split()) if split else None, **v))

exps = []
for r in wb['Expense Tracker'].iter_rows(min_row=2, values_only=True):
    if not (r[0] or r[1]) or num(r[3]) is None: continue
    exps.append(dict(date=r[0].strftime('%Y-%m-%d') if hasattr(r[0],'strftime') else None, item=r[1], cat=r[2], amt=num(r[3]), notes=r[4]))

L = []
L.append("""-- OCOTS OS: import Job Tracker + Expense Tracker history (safe to run again)
begin;
create extension if not exists pg_trgm;
alter table jobs add column if not exists customer_id uuid references customers(id);
alter table jobs add column if not exists service text;
alter table jobs add column if not exists payment_method text;
alter table jobs add column if not exists split text;
alter table jobs add column if not exists notes text;
alter table jobs add column if not exists filter_part text;
alter table jobs add column if not exists vehicle_text text;
alter table expenses add column if not exists description text;
alter table expenses add column if not exists source text;
alter table customers add column if not exists source text;
alter table vehicles add column if not exists source text;

-- clear an earlier run of this import
delete from jobs where imported_from = 'job_tracker';
delete from vehicles where source = 'job_tracker' and id not in (select vehicle_id from jobs where vehicle_id is not null);
delete from expenses where source = 'expense_tracker';

create temp table jt (n int, d date, name text, raw text, service text, price numeric, oil numeric, filt numeric, other numeric,
  notes text, method text, split text, yr int, make text, model text, engine text, grade text, cap numeric, filter text,
  customer_id uuid, score real, vehicle_id uuid) on commit drop;
insert into jt (n,d,name,raw,service,price,oil,filt,other,notes,method,split,yr,make,model,engine,grade,cap,filter) values""")
vals = []
for i,x in enumerate(rows,1):
    vals.append(f"({i},{q(x['date'])},{q(x['name'])},{q(x['raw'])},{q(x['service'])},{q(x['price'])},{q(x['oil'])},{q(x['filt'])},{q(x['other'])},{q(x['notes'])},{q(x['method'])},{q(x['split'])},{q(x['year'])},{q(x['make'])},{q(x['model'])},{q(x['engine'])},{q(x['grade'])},{q(x['cap'])},{q(x['filter'])})")
L.append(',\n'.join(vals) + ';')
L.append("""
-- match each name to a customer: best name similarity, needs 0.45+
update jt set customer_id = m.id, score = m.s from (
  select distinct on (jt.n) jt.n, c.id,
    similarity(lower(jt.name), lower(trim(coalesce(c.first_name,'') || ' ' || coalesce(c.last_name,'')))) s
  from jt cross join customers c
  order by jt.n, s desc
) m where m.n = jt.n and m.s >= 0.45;

-- one-word names (e.g. 'Karan') only match if exactly one customer has that first name
update jt set customer_id = c.id, score = 0.9 from customers c
where jt.customer_id is null and jt.name !~ ' '
  and lower(trim(c.first_name)) = lower(jt.name)
  and (select count(*) from customers c2 where lower(trim(c2.first_name)) = lower(jt.name)) = 1;

-- anyone still unmatched becomes a new customer (no phone), so no history is lost
insert into customers (first_name, last_name, notes, source)
select distinct initcap(split_part(name,' ',1)), nullif(initcap(substr(name, length(split_part(name,' ',1))+2)),''),
  'Added from Job Tracker (not found in Square)', 'job_tracker'
from jt where customer_id is null
  and not exists (select 1 from customers c where c.source='job_tracker'
    and lower(trim(coalesce(c.first_name,'')||' '||coalesce(c.last_name,''))) = lower(jt.name));
update jt set customer_id = c.id, score = 0 from customers c
where jt.customer_id is null and c.source='job_tracker'
  and lower(trim(coalesce(c.first_name,'')||' '||coalesce(c.last_name,''))) = lower(jt.name);

-- vehicles: one per customer + year + make + model
insert into vehicles (customer_id, year, make, model, engine, source)
select distinct on (customer_id, yr, make, model) customer_id, yr, make, model, engine, 'job_tracker'
from jt where make is not null
  and not exists (select 1 from vehicles v where v.customer_id = jt.customer_id and v.year is not distinct from jt.yr
    and v.make = jt.make and v.model is not distinct from jt.model);
update jt set vehicle_id = v.id from vehicles v
where v.customer_id = jt.customer_id and v.year is not distinct from jt.yr and v.make = jt.make and v.model is not distinct from jt.model;

-- jobs
insert into jobs (customer_id, vehicle_id, done_at, oil_grade, oil_litres, revenue, cost_oil, cost_parts, cost_other,
  service, payment_method, split, notes, filter_part, vehicle_text, imported_from)
select customer_id, vehicle_id, (d + time '12:00') at time zone 'America/Halifax', grade, cap, price, oil, filt, other,
  service, method, split, notes, filter, raw, 'job_tracker'
from jt order by n;

-- link each job to its Square booking on the same day, when there is exactly one
update jobs j set booking_id = b.id from bookings b
where j.imported_from = 'job_tracker' and j.booking_id is null and b.customer_id = j.customer_id
  and (b.starts_at at time zone 'America/Halifax')::date = (j.done_at at time zone 'America/Halifax')::date
  and (select count(*) from bookings b2 where b2.customer_id = j.customer_id
       and (b2.starts_at at time zone 'America/Halifax')::date = (j.done_at at time zone 'America/Halifax')::date) = 1;

-- expenses
insert into expenses (spent_at, description, category, total, source, parsed) values""")
ev = []
for e in exps:
    ev.append(f"(coalesce({q(e['date'])}::date, current_date),{q(e['item'])},{q((e['cat'] or 'other').lower())},{q(e['amt'])},'expense_tracker',{q(json.dumps({'notes':e['notes']}) if e['notes'] else None)}::jsonb)")
L.append(',\n'.join(ev) + ';')
L.append("""
commit;

-- REVIEW: names that were matched loosely or added as new customers
select n as "row", d as "date", name as "job tracker name",
  case when score = 0 then 'NEW customer (not in Square)'
       else trim(coalesce(c.first_name,'')||' '||coalesce(c.last_name,'')) end as "matched to",
  round(score::numeric,2) as "match"
from jt_review r left join customers c on c.id = r.customer_id
where score < 0.75 order by score, n;
""")
sql = '\n'.join(L)
# keep a copy of the match results for the review query (jt is dropped at commit)
sql = sql.replace("\n-- vehicles:", "\ncreate table if not exists jt_review (n int, d date, name text, customer_id uuid, score real);\ntruncate jt_review;\ninsert into jt_review select n,d,name,customer_id,score from jt;\nalter table jt_review enable row level security;\n\n-- vehicles:",1)
open(out,'w').write(sql)
print(len(rows),'jobs',len(exps),'expenses')
unparsed=[r for r in rows if r['raw'] and not r['make']]
print('no make:',[(r['name'],r['raw']) for r in unparsed])
for r in rows[:200]:
    print(r['year'],r['make'],'|',r['model'],'|',r['engine'],r['grade'],r['cap'],r['filter'],'<-',(r['raw'] or '')[:60])
